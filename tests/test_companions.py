"""P1-C: companions refresh on source change, protect user edits, stub upgrades."""

from conftest import load


def _docx(path, text):
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    d.save(str(path))


def _run(tm, root, *args):
    import io
    from contextlib import redirect_stdout
    out = io.StringIO()
    with redirect_stdout(out):
        tm.main([str(root), *args])
    return out.getvalue()


def test_pptx_carries_downgrade_warning(tmp_path):
    from pptx import Presentation
    tm = load("to_markdown")
    p = Presentation()
    p.slides.add_slide(p.slide_layouts[6])
    p.save(str(tmp_path / "d.pptx"))
    _run(tm, tmp_path)
    md = (tmp_path / "d.md").read_text(encoding="utf-8")
    assert "Auto-extracted slide text" in md


def test_first_gen_then_skip_then_update_on_source_change(tmp_path):
    tm = load("to_markdown")
    _docx(tmp_path / "n.docx", "v1")
    o1 = _run(tm, tmp_path)
    assert "converted: 1" in o1
    o2 = _run(tm, tmp_path)
    assert "skipped (up to date): 1" in o2
    _docx(tmp_path / "n.docx", "TOTALLY NEW")
    o3 = _run(tm, tmp_path)
    assert "updated: 1" in o3
    assert "TOTALLY NEW" in (tmp_path / "n.md").read_text(encoding="utf-8")


def test_user_edited_md_conflicts_force_backs_up(tmp_path):
    tm = load("to_markdown")
    _docx(tmp_path / "n.docx", "v1")
    _run(tm, tmp_path)
    md = tmp_path / "n.md"
    md.write_text(md.read_text(encoding="utf-8") + "\n- my note\n", encoding="utf-8")
    _docx(tmp_path / "n.docx", "v2")
    o = _run(tm, tmp_path)
    assert "conflicts" in o and "my note" in md.read_text(encoding="utf-8")
    _run(tm, tmp_path, "--force")
    assert (tmp_path / "n.md.localbak").exists()
    assert "my note" not in md.read_text(encoding="utf-8")


def test_handwritten_same_name_kept(tmp_path):
    tm = load("to_markdown")
    _docx(tmp_path / "n.docx", "v1")
    (tmp_path / "n.md").write_text("# my own notes\n", encoding="utf-8")
    o = _run(tm, tmp_path)
    assert "user-owned" in o
    assert (tmp_path / "n.md").read_text(encoding="utf-8") == "# my own notes\n"


def test_stub_upgrades_after_install(tmp_path):
    tm = load("to_markdown")
    _docx(tmp_path / "n.docx", "v1")
    tm.HAS_DOCX = False
    try:
        o1 = _run(tm, tmp_path)
    finally:
        tm.HAS_DOCX = True
    assert "converted: 1" in o1
    assert "complete: false" in (tmp_path / "n.md").read_text(encoding="utf-8")
    o2 = _run(tm, tmp_path)
    assert "updated: 1" in o2
    assert "complete: true" in (tmp_path / "n.md").read_text(encoding="utf-8")


def test_frontmatter_quotes_special_names(tmp_path):
    tm = load("to_markdown")
    _docx(tmp_path / 'a"b\\c.docx', "v1")
    _run(tm, tmp_path)
    md = tmp_path / 'a"b\\c.md'
    assert md.exists()
    import json
    line = [ln for ln in md.read_text(encoding="utf-8").splitlines()
            if ln.startswith("source_file:")][0]
    assert json.loads(line.split(":", 1)[1].strip()) == 'a"b\\c.docx'
